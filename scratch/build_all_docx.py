import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_styled_docx(title_text, sections_data, output_path):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    NAVY = RGBColor(27, 54, 93)
    ACCENT_BLUE = RGBColor(30, 144, 255)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(12)

    def add_h1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)

    def add_h2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = ACCENT_BLUE
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.bold = True
        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(11)

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.bold = True
        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(11)

    add_title(title_text)

    for sec in sections_data:
        stype = sec.get("type")
        if stype == "h1":
            add_h1(sec["text"])
        elif stype == "h2":
            add_h2(sec["text"])
        elif stype == "body":
            add_body(sec["text"], sec.get("bold_prefix"))
        elif stype == "bullet":
            add_bullet(sec["text"], sec.get("bold_prefix"))
        elif stype == "table":
            t = doc.add_table(rows=1, cols=len(sec["headers"]))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_cells = t.rows[0].cells
            for idx, htext in enumerate(sec["headers"]):
                hdr_cells[idx].text = htext
                shd = parse_xml(r'<w:shd {} w:fill="1B365D"/>'.format(nsdecls('w')))
                hdr_cells[idx]._tc.get_or_add_tcPr().append(shd)
                for p in hdr_cells[idx].paragraphs:
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
            for row in sec["data"]:
                r_cells = t.add_row().cells
                for idx, val in enumerate(row):
                    r_cells[idx].text = str(val)

    doc.save(output_path)
    print(f"Generated DOCX at {output_path}")

# Build Plan DOCX
plan_sections = [
    {"type": "h1", "text": "1. Executive Summary & Status"},
    {"type": "body", "text": "All 4 implementation phases of QueryMind are 100% completed with a 100% green pytest pass rate (24/24 tests passing)."},
    {"type": "h1", "text": "2. Phase Execution Details"},
    {"type": "h2", "text": "Phase 1: Test Stabilization (Completed)"},
    {"type": "bullet", "bold_prefix": "Task 1.1: ", "text": "Refactored graph builder checkpointer & fixed MemorySaver requirement in unit tests."},
    {"type": "bullet", "bold_prefix": "Task 1.2: ", "text": "Installed Playwright Chromium binaries and updated test_ui.py process runner."},
    {"type": "h2", "text": "Phase 2: Benchmark Evaluation (Completed)"},
    {"type": "bullet", "bold_prefix": "Task 2.1: ", "text": "Ran BIRD Mini-Dev student_club benchmark (48 questions, 100% gold execution accuracy)."},
    {"type": "bullet", "bold_prefix": "Task 2.2: ", "text": "Executed 5-rung ablation study comparing engine baseline vs multi-agent graph."},
    {"type": "h2", "text": "Phase 3: Feature Enhancements (Completed)"},
    {"type": "bullet", "bold_prefix": "Task 3.1: ", "text": "Automated chart detection and Streamlit bar chart rendering for tabular results."},
    {"type": "bullet", "bold_prefix": "Task 3.2: ", "text": "Distinct column value sampling in Schema Linker to match categorical text tokens."},
    {"type": "bullet", "bold_prefix": "Task 3.3: ", "text": "Intelligent sub-question planning and LLM plan revision logic."},
    {"type": "h2", "text": "Phase 4: Observability & UX (Completed)"},
    {"type": "bullet", "bold_prefix": "Task 4.1: ", "text": "Machine-readable run trace logging (logs/traces/trace_*.json)."},
    {"type": "bullet", "bold_prefix": "Task 4.2: ", "text": "Streamlit UI trace file notification and flow status indicators."},
    {"type": "bullet", "bold_prefix": "Task 4.3: ", "text": "Abstract BaseDBExecutor interface for PostgreSQL/MySQL adapter readiness."},
    {"type": "h1", "text": "3. Milestone Summary"},
    {
        "type": "table",
        "headers": ["Phase", "Deliverable", "Status"],
        "data": [
            ["Phase 1: Test Fixes", "24/24 unit & UI tests green", "🟢 Completed"],
            ["Phase 2: Evaluation", "BIRD eval & ablation JSON reports", "🟢 Completed"],
            ["Phase 3: Features", "Charts, value sampling, planning", "🟢 Completed"],
            ["Phase 4: Tracing & DB", "JSON run traces & DB adapter base", "🟢 Completed"]
        ]
    }
]

create_styled_docx(
    "QueryMind — Implementation Plan & Final Status",
    plan_sections,
    "project docs/QUERYMIND_IMPLEMENTATION_PLAN.docx"
)

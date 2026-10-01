import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_demo_questions_docx():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    NAVY = RGBColor(27, 54, 93)
    ACCENT_BLUE = RGBColor(30, 144, 255)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(12)

    def add_h1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)

    def add_h2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = ACCENT_BLUE
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.bold = True
        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(11)

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.bold = True
        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(11)

    add_title("QueryMind — Multi-Agent Live Demo Questions Guide")

    add_h1("1. Executive Demo Strategy")
    add_body("This document provides a curated list of high-impact demo questions for all 12 databases in QueryMind, specifically designed to showcase the multi-agent investigation pipeline (Profiler -> Manager -> Analyst -> Critic -> Report).")

    add_h1("2. Curated Demo Questions by Database")

    db_demos = [
        ("1. Student Club (student_club.sqlite)", [
            ("Multi-Step Planning: ", "Which club events exceeded their budget allocation, and what were the total expenses incurred for those events?"),
            ("Value Sampling: ", "Among the students who attended the 'Women's Soccer' event, how many requested a Medium T-shirt size?"),
            ("Chart Showcase: ", "List total member expenses grouped by event status.")
        ]),
        ("2. Financial (financial.sqlite)", [
            ("Opaque Schema Linking: ", "What is the total loan amount for clients living in the Prague district?"),
            ("Critic Guard Safety: ", "Find accounts with a negative balance after transactions in December 1998, and list their issued credit cards."),
            ("Chart Showcase: ", "Show total transaction amounts per transaction type in 1997.")
        ]),
        ("3. Formula 1 (formula_1.sqlite)", [
            ("Multi-Table Join: ", "Which driver had the fastest pit stop duration in the 2021 Monaco Grand Prix?"),
            ("Comparative Planning: ", "Compare the total race wins between Ferrari and Mercedes constructors."),
            ("Standings Chart: ", "List the top 5 drivers by total career race wins.")
        ]),
        ("4. European Football (european_football_2.sqlite)", [
            ("Complex Filtering: ", "Which team scored the highest total goals in home matches during the 2015 season?"),
            ("Goal Distribution Chart: ", "Show total goals scored per league in the 2014 season.")
        ]),
        ("5. California Schools (california_schools.sqlite)", [
            ("Demographic Analysis: ", "Which school district has the highest average SAT math score?"),
            ("Auto-Repair Loop: ", "What is the correlation between free lunch eligible students and average reading scores?")
        ]),
        ("6. Superhero (superhero.sqlite)", [
            ("Multi-Condition Join: ", "How many male superheroes published by Marvel Comics have super strength and flight powers?"),
            ("Publisher Breakdown Chart: ", "Show the count of superheroes grouped by publisher and alignment (good, bad, neutral).")
        ]),
        ("7. Card Games (card_games.sqlite)", [
            ("Category Filtering: ", "What is the average power of creature cards in the 'Rare' category?"),
            ("Set Distribution Chart: ", "List the total number of cards released in each card set.")
        ]),
        ("8. Thrombosis Prediction (thrombosis_prediction.sqlite)", [
            ("Clinical Planning: ", "What is the average age of patients diagnosed with thrombosis, and what are their most common lab test abnormalities?"),
            ("Critic Guard Anti-Hallucination: ", "Find patients under age 20 with severe thrombosis complications.")
        ]),
        ("9. Toxicology (toxicology.sqlite)", [
            ("Scientific Property Lookup: ", "How many chemical compounds with a molecular weight greater than 300 passed the toxicity assay test?")
        ]),
        ("10. Debit Card Specializing (debit_card_specializing.sqlite)", [
            ("Date Expression Analysis: ", "What was the average monthly fuel consumption of SME business customers in 2013?")
        ]),
        ("11. Codebase Community (codebase_community.sqlite)", [
            ("User Tag Join: ", "Which users have the highest reputation among those who answered questions tagged with 'Python'?")
        ]),
        ("12. Default Sales Fixture (fixture.sqlite)", [
            ("Revenue Summary: ", "What is the total revenue across all sales?"),
            ("Category Chart: ", "List total sales amount grouped by product category.")
        ])
    ]

    for db_title, questions in db_demos:
        add_h2(db_title)
        for prefix, q_text in questions:
            add_bullet(q_text, prefix)

    add_h1("3. Recommended 5-Minute Presentation Flow")
    add_bullet("Select 'BIRD: Student Club' from the Streamlit sidebar dropdown.", "Step 1: ")
    add_bullet("Ask: 'Which club events exceeded their budget allocation, and what were the total expenses incurred for those events?'", "Step 2: ")
    add_bullet("Highlight the Manager Agent's investigation plan during the Human-in-the-Loop pause.", "Step 3: ")
    add_bullet("Click 'Approve & Execute queries'.", "Step 4: ")
    add_bullet("Showcase the Profiler row counts, Analyst SQL, Critic verdict, grounded report, and Streamlit bar chart.", "Step 5: ")

    output_path = "project docs/QUERYMIND_DEMO_QUESTIONS.docx"
    doc.save(output_path)
    print(f"Generated DOCX at {output_path}")

if __name__ == "__main__":
    build_demo_questions_docx()

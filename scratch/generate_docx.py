import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Set standard 1 inch margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Primary color palette
NAVY = RGBColor(27, 54, 93)
DARK_GRAY = RGBColor(60, 60, 60)
ACCENT_BLUE = RGBColor(30, 144, 255)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(12)

def add_h1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = NAVY
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)

def add_h2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = ACCENT_BLUE
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)

def add_body(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(11)

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(11)

# Document Content Generation
add_title("QueryMind — Comprehensive Project Analysis & Architecture")

add_h1("1. Executive Summary & Project Purpose")
add_body("QueryMind is an open-source, multi-agent Autonomous Data Analyst system built in Python (v3.11+). It leverages LangGraph, SQLite, and Google Gemini (or other provider-swappable LLMs) to investigate complex business questions over relational databases.")

add_h2("Core Objectives")
add_bullet(" Answer natural-language business questions by decomposing them into sub-questions, executing SQL queries, validating evidence, and synthesizing grounded reports.", "1. End-to-End Investigation:")
add_bullet(" Decouple text-to-SQL logic (SchemaLinker, SQLEngine, DBExecutor) so it can be evaluated independently on standard Text-to-SQL benchmarks (e.g., BIRD Mini-Dev).", "2. Benchmarkable Inner Engine:")
add_bullet(" Employ a cyclic loop (Manager → Analyst → Critic → Report) with a hard critic guard to re-investigate when initial evidence is missing or insufficient.", "3. Agentic Planning & Self-Critique:")
add_bullet(" Designed to run efficiently on standard hardware without local GPUs using free-tier LLMs, backed by local caching (.llm_cache.sqlite) and vector search via ChromaDB.", "4. $0 / Low-Cost Operation:")

add_h1("2. Architecture & System Design (Simplified & Intuitive)")
add_body("To make QueryMind's architecture easy to understand, think of the system as an Autonomous Corporate Analytics Team working in a synchronized pipeline:")

# Table representing team roles
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table.rows[0].cells
headers = ["Role / Agent", "Analogy", "Key Responsibilities"]
for idx, text in enumerate(headers):
    hdr_cells[idx].text = text
    shading_xml = parse_xml(r'<w:shd {} w:fill="1B365D"/>'.format(nsdecls('w')))
    hdr_cells[idx]._tc.get_or_add_tcPr().append(shading_xml)
    for p in hdr_cells[idx].paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

roles_data = [
    ("👤 User / Executive", "The Client", "Submits high-level analytical business questions (e.g., 'Why did Q3 revenue drop?')."),
    ("🔍 Data Profiler", "Data Clerk", "Inspects matched tables, counts database rows, and provides immediate schema context."),
    ("📝 Manager Agent", "Team Lead", "Breaks complex questions into targeted SQL sub-questions and revises plans if evidence is weak."),
    ("🛠️ Analyst Agent", "Data Engineer", "Runs schema linking (BM25 + ChromaDB), generates dialect-correct SQL, and executes auto-repair loops on errors."),
    ("⚖️ Critic Guard", "QA / Auditor", "Deterministic safety officer. Enforces that empty or failed queries CANNOT be marked as sufficient."),
    ("📊 Report Agent", "Business Writer", "Synthesizes final natural-language answers strictly grounded on returned query numbers.")
]

for r_name, r_analogy, r_resp in roles_data:
    row_cells = table.add_row().cells
    row_cells[0].text = r_name
    row_cells[1].text = r_analogy
    row_cells[2].text = r_resp

add_h2("System Layer Overview")
table_layers = doc.add_table(rows=1, cols=3)
table_layers.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells2 = table_layers.rows[0].cells
for idx, text in enumerate(["Layer", "Primary Component", "Function"]):
    hdr_cells2[idx].text = text
    shading_xml = parse_xml(r'<w:shd {} w:fill="1B365D"/>'.format(nsdecls('w')))
    hdr_cells2[idx]._tc.get_or_add_tcPr().append(shading_xml)
    for p in hdr_cells2[idx].paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

layers_data = [
    ("Interface Layer", "app.py (Streamlit)", "Provides live step-by-step trace visualization and Human-in-the-Loop (HITL) plan approval."),
    ("Orchestration Layer", "graph/builder.py (LangGraph)", "Manages state transitions, iteration limits, cyclic fallback loops, and graph checkpoints."),
    ("Text-to-SQL Engine", "engine/ (sql_gen, schema, db)", "Independent module for schema retrieval, prompt construction, read-only SQL execution & repair."),
    ("LLM & Cache Layer", "llm/client.py (Gemini + Caching)", "Provider-swappable client with 429 rate limit retries and prompt hashing to .llm_cache.sqlite.")
]

for l_name, l_comp, l_func in layers_data:
    row_cells = table_layers.add_row().cells
    row_cells[0].text = l_name
    row_cells[1].text = l_comp
    row_cells[2].text = l_func

add_h1("3. Core Modules & Code Walkthrough")
add_bullet(" Inspects sqlite_master and PRAGMA table_info to extract schemas and DDL.", "catalog.py:")
add_bullet(" Implements hybrid BM25 token overlap + ChromaDB semantic vector schema linking.", "schema.py:")
add_bullet(" Enforces read-only SQLite execution with URI mode (?mode=ro), timeouts, and row limits.", "db.py:")
add_bullet(" AST/Regex guard blocking DDL/DML mutation keywords (INSERT, DROP, UPDATE).", "sql_safety.py:")
add_bullet(" Translates questions to SQL and runs an iterative error repair loop on failures.", "sql_gen.py:")

add_h1("4. Benchmark & Test Suite Health")
add_body("Running pytest on the repository yields 18 passing tests out of 21. Two graph tests require thread identifier configuration due to default MemorySaver checkpointers, and the UI test requires Playwright browser binaries.")

add_h1("5. Recommendations & Next Steps")
add_bullet(" Pass configurable thread_id in test_graph.py to bring all graph tests to green.", "1. Fix Graph Tests:")
add_bullet(" Install browser drivers via 'playwright install' to enable automated end-to-end UI testing.", "2. Setup Playwright:")
add_bullet(" Execute 'python -m querymind.eval --db-id student_club' to benchmark baseline execution accuracy.", "3. Benchmark Run:")

output_path = "project docs/QUERYMIND_ANALYSIS.docx"
doc.save(output_path)
print(f"Successfully generated DOCX at {output_path}")
